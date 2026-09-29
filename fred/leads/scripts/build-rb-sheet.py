#!/usr/bin/env python3
"""Build the Richards Bay no-website lead spreadsheet (xlsx -> Google Sheet).

Source: fred/products/leads/gmaps_*.csv (Google Maps scan 2026-09-28).
Output: fred/leads/out/Richards-Bay-leads-2026-09-28.xlsx
"""
import csv
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = "/Users/deonvandenberg/.openclaw/workspace/fred/products/leads"
OUT = "/Users/deonvandenberg/.openclaw/workspace/fred/leads/out/Richards-Bay-leads-2026-09-28.xlsx"

GOLD = "C8A34E"
GOLD_DEEP = "A8863A"
CREAM = "FAF8F3"
INK = "14142A"

HDR_FONT = Font(name="Outfit", bold=True, color="FFFFFF", size=11)
HDR_FILL = PatternFill("solid", fgColor=GOLD_DEEP)
BODY_FONT = Font(name="Outfit", size=10, color=INK)
TITLE_FONT = Font(name="Outfit", bold=True, size=14, color=INK)
LABEL_FONT = Font(name="Outfit", bold=True, size=10, color=GOLD_DEEP)
THIN = Side(style="thin", color="EADFC4")
BORDER = Border(bottom=THIN)


def load(name):
    with open(f"{BASE}/{name}", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


allb = load("gmaps_all_businesses.csv")
noweb = load("gmaps_no_website.csv")
contactable = load("gmaps_no_website_contactable.csv")


def industry(row):
    return row.get("category", "").strip()


def phone(row):
    return row.get("phone", "").strip()


wb = Workbook()

# ---------------------------------------------------------------- Summary
ws = wb.active
ws.title = "Summary"
ws.sheet_view.showGridLines = False
ws["A1"] = "Richards Bay — businesses with no website"
ws["A1"].font = TITLE_FONT
ws["A2"] = "Google Maps scan · 2026-09-28 · AutoEffortless research"
ws["A2"].font = Font(name="Outfit", size=10, italic=True, color="6B6B6B")

stats = [
    ("Businesses scanned", len(allb)),
    ("No website", len(noweb)),
    ("No website + phone (contactable)", len(contactable)),
    ("Industries represented", len({industry(r) for r in noweb if industry(r)})),
]
r = 4
ws.cell(r, 1, "Headline").font = LABEL_FONT
r += 1
for label, val in stats:
    ws.cell(r, 1, label).font = BODY_FONT
    c = ws.cell(r, 2, val)
    c.font = Font(name="Outfit", bold=True, size=11, color=GOLD_DEEP)
    r += 1

r += 1
ws.cell(r, 1, "Top industries (by no-website count)").font = LABEL_FONT
r += 1
ws.cell(r, 1, "Industry").font = HDR_FONT
ws.cell(r, 1).fill = HDR_FILL
ws.cell(r, 2, "Businesses").font = HDR_FONT
ws.cell(r, 2).fill = HDR_FILL
r += 1
counts = Counter(industry(x) or "Uncategorised" for x in noweb)
for name, n in counts.most_common():
    ws.cell(r, 1, name).font = BODY_FONT
    c = ws.cell(r, 2, n)
    c.font = BODY_FONT
    r += 1

ws.column_dimensions["A"].width = 46
ws.column_dimensions["B"].width = 14


# ---------------------------------------------------------------- Lists
HEADERS = ["#", "Business", "Industry", "Phone", "Address", "Suburb",
           "Rating", "Reviews", "Status", "Maps link"]
WIDTHS = [5, 38, 24, 16, 52, 18, 8, 9, 10, 20]


def add_sheet(title, rows):
    s = wb.create_sheet(title)
    s.append(HEADERS)
    for i, cell in enumerate(s[1], start=1):
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center")
        s.column_dimensions[get_column_letter(i)].width = WIDTHS[i - 1]
    for i, row in enumerate(rows, start=1):
        rating = row.get("rating", "").strip()
        reviews = row.get("reviews_n", "").strip()
        s.append([
            i,
            row.get("name", "").strip(),
            industry(row),
            phone(row),
            row.get("address", "").strip(),
            row.get("suburb", "").strip(),
            float(rating) if rating else None,
            int(reviews) if reviews else None,
            row.get("status", "").strip(),
            row.get("place_url", "").strip(),
        ])
    for row in s.iter_rows(min_row=2, max_row=s.max_row, max_col=len(HEADERS)):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            if cell.column == 4:  # phone — keep as text
                cell.number_format = "@"
            if cell.column == 10:  # maps link
                v = cell.value
                if v:
                    cell.hyperlink = v
                    cell.font = Font(name="Outfit", size=10, color="1155CC", underline="single")
    s.freeze_panes = "A2"
    s.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{s.max_row}"
    return s


add_sheet("No website + phone (660)", contactable)
add_sheet("No website (711)", noweb)
add_sheet("All scanned (1409)", allb)

wb.save(OUT)
print("wrote", OUT)
